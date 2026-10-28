# -*- coding: utf-8 -*-
"""Derive les cellules « Non concerne » et leur cause datee.

Usage :
    python3 scripts/non_concerne.py \
        --comparateur COMPARATEUR_V2.xlsx \
        --fenetres donnees/fenetres.json \
        --sortie sortie/non_concerne.csv

Principe : un candidat qui ne siegeait ni a l'Assemblee ni au Senat le jour du
scrutin ne pouvait pas voter. Ce n'est pas une absence, c'est une impossibilite,
et elle se deduit des dates de mandat. Le script ecrit la valeur et la cause,
avec la date dans la phrase pour que le lecteur puisse verifier.

Ce qu'il ne fait pas, volontairement :
  - il ne touche a aucune cellule deja remplie. Il signale les divergences et
    laisse la redaction trancher ;
  - il ne derive rien pour un candidat dont les fenetres ne sont pas bornees ;
  - il ne derive rien pour un texte sans date exploitable.
"""
import argparse
import csv
import json
import re
from datetime import date
from pathlib import Path

import openpyxl

MOIS = [u"janvier", u"février", u"mars", u"avril", u"mai", u"juin", u"juillet",
        u"août", u"septembre", u"octobre", u"novembre", u"décembre"]
JJMMAAAA = re.compile(r"\b(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\b")


def propre(v):
    return str(v).replace("\r\n", "\n").strip() if v is not None else ""


def en_lettres(d):
    return u"%d%s %s %d" % (d.day, u"er" if d.day == 1 else u"", MOIS[d.month - 1], d.year)


def lit_date(texte):
    """La premiere date jj/mm/aaaa d'une cellule de date redigee en prose."""
    m = JJMMAAAA.search(propre(texte))
    if not m:
        return None
    j, mo, a = (int(x) for x in m.groups())
    try:
        return date(a, mo, j)
    except ValueError:
        return None


def iso(s):
    return date(*(int(x) for x in s.split("-"))) if s else None


def couvre(fiche, jour):
    """La ou les chambres ou le candidat siegeait ce jour-la, et l'empechement
    eventuel. Rend (chambres, motif_empechement)."""
    for e in fiche["empeches"]:
        d, f = iso(e["debut"]), iso(e["fin"])
        if d <= jour <= f:
            return [], e
    chambres = []
    for f in fiche["fenetres"]:
        d = iso(f["debut"])
        fin = iso(f["fin"])
        if d <= jour and (fin is None or jour <= fin):
            chambres.append(f["chambre"])
    return chambres, None


def cause(fiche, jour, empeche):
    jl = en_lettres(jour)
    if empeche:
        return (u"Non concerné : %s du %s au %s, il ne pouvait pas voter le %s."
                % (empeche["motif"], en_lettres(iso(empeche["debut"])),
                   en_lettres(iso(empeche["fin"])), jl))
    if fiche["etat"] == "jamais":
        return u"Non concerné : %s Aucun vote possible le %s." % (fiche["note"], jl)
    return (u"Non concerné : ne siégeait ni à l'Assemblée nationale ni au Sénat le %s." % jl)


def textes(chemin):
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    ws = wb["1_TEXTES"]
    it = ws.iter_rows(values_only=True)
    e = [propre(c) for c in next(it)]
    i = {k: n for n, k in enumerate(e)}
    out = []
    for r in it:
        if not r or not r[0]:
            continue
        g = lambda k: propre(r[i[k]]) if k in i and i[k] < len(r) else ""
        out.append({"id": g("Texte_ID"), "nom": g("Nom de la fiche").replace("\n", " "),
                    "etat": g(u"État"), "brut": g("Date"), "jour": lit_date(g("Date"))})
    wb.close()
    return out


def positions(chemin):
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    ws = wb["2_POSITIONS"]
    it = ws.iter_rows(values_only=True)
    e = [propre(c) for c in next(it)]
    i = {k: n for n, k in enumerate(e)}
    out = {}
    for r in it:
        if not r or not r[0]:
            continue
        g = lambda k: propre(r[i[k]]) if k in i and i[k] < len(r) else ""
        out[(g("Texte_ID"), g("Candidat"))] = (g("Vote personnel"), g(u"Cause si non concerné"), g(u"État"))
    wb.close()
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--comparateur", required=True)
    ap.add_argument("--fenetres", required=True)
    ap.add_argument("--sortie", required=True)
    a = ap.parse_args()

    ref = json.loads(Path(a.fenetres).read_text(encoding="utf-8"))["candidats"]
    tx = textes(a.comparateur)
    pos = positions(a.comparateur)

    sans_date = [t for t in tx if not t["jour"]]
    a_borner = sorted(n for n, f in ref.items() if f["etat"] == "a_borner")

    lignes, divergences = [], []
    pouvait = 0
    for t in tx:
        if not t["jour"]:
            continue
        for nom, fiche in sorted(ref.items()):
            if fiche["etat"] == "a_borner":
                continue
            chambres, empeche = couvre(fiche, t["jour"])
            if chambres:
                pouvait += 1
                continue
            c = cause(fiche, t["jour"], empeche)
            actuel, cause_actuelle, etat = pos.get((t["id"], nom), ("", "", ""))
            if actuel and actuel != u"Non concerné":
                divergences.append((t["id"], nom, actuel, c))
                continue
            lignes.append({
                "Texte_ID": t["id"], "Candidat": nom,
                "Vote personnel": u"Non concerné", u"Cause si non concerné": c,
                "statut": u"déjà rempli" if actuel else u"à remplir",
                "cause actuelle": cause_actuelle,
            })

    Path(a.sortie).parent.mkdir(parents=True, exist_ok=True)
    with open(a.sortie, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["Texte_ID", "Candidat", "Vote personnel",
                                          u"Cause si non concerné", "statut", "cause actuelle"],
                           delimiter=";")
        w.writeheader()
        for l in lignes:
            w.writerow(l)

    neuf = sum(1 for l in lignes if l["statut"] == u"à remplir")
    print(u"textes : %d  (dont %d sans date exploitable)" % (len(tx), len(sans_date)))
    print(u"candidats derivables : %d sur %d" % (len(ref) - len(a_borner), len(ref)))
    print(u"cellules « non concerné » derivees : %d" % len(lignes))
    print(u"   dont nouvelles : %d" % neuf)
    print(u"   dont confirmant une cellule existante : %d" % (len(lignes) - neuf))
    print(u"cellules ou le candidat pouvait voter : %d  (a documenter a la main)" % pouvait)
    if a_borner:
        print(u"\nfenetres a borner avant de deriver : %s" % ", ".join(a_borner))
    if sans_date:
        print(u"\ntextes sans date exploitable : %d" % len(sans_date))
        for t in sans_date[:12]:
            print(u"   %-10s %-44s %r" % (t["id"], t["nom"][:44], t["brut"][:40]))
    if divergences:
        print(u"\nDIVERGENCES : la machine dit « non concerné », le tableau dit autre chose")
        for tid, nom, actuel, c in divergences:
            print(u"   %-10s %-24s tableau : %s" % (tid, nom, actuel))
    print(u"\nfichier : %s" % a.sortie)


if __name__ == "__main__":
    main()
